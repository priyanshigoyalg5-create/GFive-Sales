#!/usr/bin/env python3
"""Rebuild app/seed.json from the owner's ERP export (DETAILS.xlsx: sheets CREDITOR, JOBBER, ITEM DETAILS)
and embed it into app/index.html (the app reads the const SEED inside index.html, NOT the seed.json file).
Usage: python3 tools/build_seed.py DETAILS.xlsx app/seed.json app/index.html [Abdul_Kalam_Job_Work_Register.xlsx]
The optional 4th file is the old Excel register: its Fabric Master (F numbers missing in the ERP export) and design numbers are merged in.
DETAILS.xlsx holds real vendor GST numbers and phones: keep it private, never commit it publicly."""
import re,sys,json
from openpyxl import load_workbook
xlsx,seedp,htmlp=sys.argv[1:4]
reg=sys.argv[4] if len(sys.argv)>4 else None
wb=load_workbook(xlsx,data_only=True)
def ph(v):
    if v is None: return None
    s=str(v).strip(); s=s[:-2] if s.endswith('.0') else s
    return s or None
def cl(v):
    if v is None: return None
    s=re.sub(r'\s+',' ',str(v).replace('\xa0',' ')).strip(); return s or None
ven=[]
for r in wb['CREDITOR'].iter_rows(min_row=2,values_only=True):
    if not r[1]: continue
    addr=", ".join([x for x in [cl(r[5]),cl(r[6]),cl(r[8])] if x])
    if addr and r[7]:
        try: addr+=" - "+str(int(r[7]))
        except: pass
    o={"name":cl(r[1])}
    if cl(r[3]): o["gstin"]=cl(r[3]).upper()
    if addr: o["address"]=addr
    p=ph(r[14]) or ph(r[16])
    if p: o["phone"]=p
    ven.append(o)
RULES=[("KATHA","Katha"),("DYER","Dyeing"),("DAYER","Dyeing"),("DYEING","Dyeing"),("HAND","Hand work"),("TAILOR","Tailor / Finishing"),("STITCH","Tailor / Finishing"),("FINISH","Tailor / Finishing"),("MACHINE","Machine Embroidery"),("COMPUTER","Computer work"),("-COM","Computer work"),("KHAKHA","Khaka"),("KHAKA","Khaka"),("PRINT","Print"),("APPLI","Applique work"),("APLIC","Applique work")]
kar=[]
for r in wb['JOBBER'].iter_rows(min_row=2,values_only=True):
    if not r[1]: continue
    name=cl(r[1]); o={"name":name}
    for t,c in RULES:
        if t in name.upper(): o["proc"]=c; break
    p=ph(r[14]) or ph(r[16])
    if p: o["phone"]=p
    kar.append(o)
rx=re.compile(r'(?<![A-Z0-9])F[- ]?(\d{1,4})(?!\d)')
fab={}; des=set()
for r in wb['ITEM DETAILS'].iter_rows(min_row=2,values_only=True):
    if not r[1]: continue
    n=cl(r[1]); m=rx.search(n); grp=str(r[8] or '')
    if m and grp.startswith('FABRIC'):            # only items that carry an F number are fabrics
        no="F"+m.group(1)
        nm=(cl(re.sub(r'\(?\bF[- ]?\d{1,4}\b\)?','',n).replace('FABRIC PCS/','')) or n).strip(' /-()') or n
        if nm.count('(')>nm.count(')'): nm=nm+')' if re.search(r'\([A-Z]$',nm) else nm.rsplit('(',1)[0].strip()
        o={"fabNo":no,"name":nm,"unit":"Mtr" if str(r[4]).upper().startswith("M") else "Pcs"}
        if r[5]: o["rate"]=r[5]
        fab.setdefault(no,o)
    d=re.sub(r'^FABRIC PCS/','',n).replace(' ','')
    if re.fullmatch(r'(SD|TY|RZ|MH|CR|GF|GBPL|GFIVE|HSN)?-?\d{3,6}[A-Z]?',d) and re.match(r'[A-Z]',d):
        des.add(d.replace('-',''))
if reg:                                              # merge the old register (does not override ERP rows)
    w2=load_workbook(reg,data_only=True)
    for r in w2['Fabric Master'].iter_rows(min_row=5,values_only=True):
        if r[1] and str(r[1]).strip() not in fab:
            o={"fabNo":str(r[1]).strip(),"name":r[2]}
            if r[3] is not None: o["rate"]=r[3]
            if r[0]: o["width"]=r[0]
            fab[o["fabNo"]]=o
    for r in w2['Issue Register'].iter_rows(min_row=5,values_only=True):
        if r[0] and r[8]: des.add(str(r[8]).strip())
cats=["Dyeing","Hand work","Katha","Tailor / Finishing","Computer work","Machine Embroidery","Applique work","Print","Khaka"]
alt=[("Dyeing (re-dye)","Dyeing"),("Hand work repair","Hand work"),("Piping repair","Tailor / Finishing"),("Stitching repair","Tailor / Finishing"),("Other","")]
seed={"vendors":ven,"karigars":kar,"fabrics":sorted(fab.values(),key=lambda x:x['fabNo']),"designs":[{"design":d} for d in sorted(des)],"categories":[{"name":c,"order":i+1} for i,c in enumerate(cats)],"alttypes":[dict(name=n,order=i+1,**({"proc":p} if p else {})) for i,(n,p) in enumerate(alt)]}
json.dump(seed,open(seedp,'w'),indent=0)
blob=json.dumps(seed,separators=(',',':'),ensure_ascii=False).replace('</','<\\/')
h=open(htmlp).read()
h2,n=re.subn(r'const SEED=\{.*?\};\n(?=async function seed\(\))',lambda m:'const SEED='+blob+';\n',h,flags=re.S)
assert n==1,'SEED block not found'
open(htmlp,'w').write(h2)
print({k:len(v) for k,v in seed.items()})
